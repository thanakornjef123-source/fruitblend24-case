"""Writes results.pkl (from 03_analysis.py) into FruitBlend24_analysis.xlsx.
P&L sheets keep live formulas (GP, operating profit, variance) so the numbers can be audited in Excel."""
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUT = Path(__file__).parent
R = pd.read_pickle(OUT / "results.pkl")
F = "Arial"
HDR = PatternFill("solid", fgColor="1F4E78")
TOT = PatternFill("solid", fgColor="F2F2F2")
THB = '#,##0;(#,##0);"-"'
THB2 = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'

wb = Workbook()
wb.remove(wb.active)


def header(ws, cols, row=1):
    for j, c in enumerate(cols, 1):
        x = ws.cell(row, j, c)
        x.font = Font(name=F, bold=True, color="FFFFFF")
        x.fill = HDR
        x.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = ws.cell(row + 1, 2)


def fmt_col(name):
    n = name.lower()
    if any(k in n for k in ["rate", "pct", "margin", "lift", "share", "ratio", "wape", "bias", "uplift", "price_change"]) and "thb" not in n and "per_cup" not in n:
        return PCT
    if "per_cup" in n or "price" in n or "elasticity" in n or n in ("z",):
        return THB2
    return THB


def write_df(name, df, index=True, note=None):
    ws = wb.create_sheet(name)
    d = df.reset_index() if index else df.copy()
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(name=F, italic=True, color="595959")
        r0 = 3
    header(ws, list(d.columns), r0)
    for i, row in enumerate(d.itertuples(index=False), r0 + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, pd.Timestamp):
                v = v.to_pydatetime()
            elif isinstance(v, float) and pd.isna(v):
                v = None
            c = ws.cell(i, j, v)
            c.font = Font(name=F)
            col = d.columns[j - 1]
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                c.number_format = fmt_col(str(col))
            elif hasattr(v, "year"):
                c.number_format = "yyyy-mm-dd"
    for j, col in enumerate(d.columns, 1):
        ws.column_dimensions[get_column_letter(j)].width = max(12, min(40, len(str(col)) + 3))
    return ws


# ------------------------------------------------------------------ README
ws = wb.create_sheet("README")
lines = [
    ("FruitBlend24 — Analysis workbook (tasks 2–5)", True),
    ("Built by 03_analysis.py from the cleaned data (FruitBlend24_clean.xlsx). All THB.", False),
    ("", False),
    ("Definitions", True),
    ("Gross profit (GP) = revenue − platform commission − fruit − packaging − labor − waste cost (fruit + packaging of discarded cups). Matches the budget's ~26% GP margin.", False),
    ("Operating profit = GP − fixed kitchen overhead (670k / month in total).", False),
    ("Blue cells in pnl_month / pnl_kitchen are inputs; black are formulas (GP, operating profit, variance).", False),
    ("", False),
    ("Sheets", True),
    ("pnl_month — monthly P&L vs budget (formulas)", False),
    ("pnl_kitchen — 12-month P&L by kitchen incl. overhead (formulas)", False),
    ("pnl_sku, pnl_kitchen_month, pnl_sku_month, pnl_kitchen_sku — P&L cuts", False),
    ("season — GP and operating profit by season vs budget", False),
    ("fruit_shock — Jun–Aug vs Mar–May fruit cost per cup and the extra cost", False),
    ("promo_eval, promo_eval_sku — promo lift vs matched baseline, elasticity, profit impact, break-even lift", False),
    ("unit_economics — margin per cup at recent costs; price_sim — GP impact of price changes under 3 elasticity cases", False),
    ("platform — LINE MAN vs Grab margin per cup; demand_hour / demand_dow — demand shape", False),
    ("forecast_backtest — 4-week same-weekday forecast accuracy, Jun–Aug 2026", False),
    ("newsvendor — critical ratio and prep buffer per SKU", False),
    ("outlook_* — Sep–Nov 2026 monthly P&L by scenario; bridge — status quo → plan, lever by lever", False),
    ("inventory_month / inventory_weekly — cups to prep and fruit purchase by month, and by week × kitchen × SKU", False),
    ("forecast_daily_plan — daily forecast per kitchen × SKU under the plan", False),
    ("backcast — the model replayed over the last 12 months (validation + full-year value of the plan)", False),
    ("assumptions — every assumption behind the numbers", False),
]
for i, (t, b) in enumerate(lines, 1):
    c = ws.cell(i, 1, t)
    c.font = Font(name=F, bold=b, size=14 if i == 1 else 10)
ws.column_dimensions["A"].width = 140

# ------------------------------------------------------------------ P&L month with formulas
pm = R["pnl_month"]
ws = wb.create_sheet("pnl_month")
cols = ["month", "units_sold", "gross_revenue_thb", "discount_given_thb", "commission_thb", "fruit_thb", "packaging_thb",
        "labor_thb", "waste_thb", "gross_profit_thb", "gp_margin", "budget_revenue_thb", "budget_gp_thb",
        "rev_var_thb", "rev_var_pct", "gp_var_thb", "gp_var_pct", "overhead_thb", "operating_profit_thb"]
header(ws, cols)
BLUE = Font(name=F, color="0000FF")
for i, (m, r) in enumerate(pm.iterrows(), 2):
    vals = [m.to_pydatetime(), r.units_sold, r.gross_revenue_thb, r.discount_thb, r.commission_thb, r.fruit_thb,
            r.packaging_thb, r.labor_thb, r.waste_thb]
    for j, v in enumerate(vals, 1):
        c = ws.cell(i, j, v)
        c.font = BLUE if j > 1 else Font(name=F)
        c.number_format = "mmm yyyy" if j == 1 else THB
    f = {10: f"=C{i}-E{i}-F{i}-G{i}-H{i}-I{i}", 11: f"=IF(C{i}=0,0,J{i}/C{i})", 12: r.budget_revenue_thb,
         13: r.budget_gp_thb, 14: f"=C{i}-L{i}", 15: f"=IF(L{i}=0,0,N{i}/L{i})", 16: f"=J{i}-M{i}",
         17: f"=IF(M{i}=0,0,P{i}/M{i})", 18: r.overhead_thb, 19: f"=J{i}-R{i}"}
    for j, v in f.items():
        c = ws.cell(i, j, v)
        c.font = BLUE if j in (12, 13, 18) else Font(name=F)
        c.number_format = PCT if j in (11, 15, 17) else THB
t = len(pm) + 2
ws.cell(t, 1, "Total").font = Font(name=F, bold=True)
for j in range(2, 20):
    L = get_column_letter(j)
    if j == 11:
        v = f"=IF(C{t}=0,0,J{t}/C{t})"
    elif j == 15:
        v = f"=IF(L{t}=0,0,N{t}/L{t})"
    elif j == 17:
        v = f"=IF(M{t}=0,0,P{t}/M{t})"
    else:
        v = f"=SUM({L}2:{L}{t-1})"
    c = ws.cell(t, j, v)
    c.font = Font(name=F, bold=True)
    c.fill = TOT
    c.number_format = PCT if j in (11, 15, 17) else THB
ws.cell(t + 2, 1, "Overhead = sum of kitchen_overhead (180k + 150k + 200k + 140k). Budget = monthly_budget sheet (flat, set by management).").font = Font(name=F, italic=True, color="595959")
for j in range(1, 20):
    ws.column_dimensions[get_column_letter(j)].width = 15

# ------------------------------------------------------------------ P&L kitchen with formulas
pk = R["pnl_kitchen"]
ws = wb.create_sheet("pnl_kitchen")
cols = ["kitchen", "units_sold", "gross_revenue_thb", "commission_thb", "fruit_thb", "packaging_thb", "labor_thb",
        "waste_thb", "gross_profit_thb", "gp_per_cup_thb", "waste_rate", "overhead_12m_thb", "operating_profit_thb", "gp_uplift_to_breakeven"]
header(ws, cols)
for i, (k, r) in enumerate(pk.iterrows(), 2):
    vals = [k, r.units_sold, r.gross_revenue_thb, r.commission_thb, r.fruit_thb, r.packaging_thb, r.labor_thb, r.waste_thb]
    for j, v in enumerate(vals, 1):
        c = ws.cell(i, j, v)
        c.font = BLUE if j > 1 else Font(name=F)
        c.number_format = THB
    for j, v, nf in [(9, f"=C{i}-D{i}-E{i}-F{i}-G{i}-H{i}", THB), (10, f"=IF(B{i}=0,0,I{i}/B{i})", THB2),
                     (11, r.waste_rate, PCT), (12, r.overhead_thb, THB), (13, f"=I{i}-L{i}", THB),
                     (14, f"=IF(I{i}=0,0,L{i}/I{i}-1)", PCT)]:
        c = ws.cell(i, j, v)
        c.font = BLUE if j in (11, 12) else Font(name=F)
        c.number_format = nf
t = len(pk) + 2
ws.cell(t, 1, "Total").font = Font(name=F, bold=True)
for j in [2, 3, 4, 5, 6, 7, 8, 9, 12, 13]:
    L = get_column_letter(j)
    c = ws.cell(t, j, f"=SUM({L}2:{L}{t-1})")
    c.font = Font(name=F, bold=True)
    c.fill = TOT
    c.number_format = THB
for j in range(1, 15):
    ws.column_dimensions[get_column_letter(j)].width = 16

# ------------------------------------------------------------------ other tables
write_df("pnl_sku", R["pnl_sku"])
write_df("pnl_kitchen_month", R["pnl_kitchen_month"])
write_df("pnl_sku_month", R["pnl_sku_month"])
write_df("pnl_kitchen_sku", R["pnl_kitchen_sku"])
write_df("season", R["season"])
write_df("fruit_shock", R["fruit_shock_JunAug"], note="Extra fruit cost Jun–Aug 2026 had cost per cup stayed at the Mar–May average")
write_df("promo_eval", R["promo_code"], note=f"Baseline = same kitchen×SKU non-promo days ±28 days, same weekday/weekend type. Other SKUs on RC101 days: {R['rc101_other_sku_change']:+.1%} (no cannibalisation). RC102 needs {R['rc102_repeat_cups_needed']:,.0f} future full-price cups to pay back ({R['rc102_repeat_per_incr_cup']:.2f} per incremental cup).")
write_df("promo_eval_sku", R["promo_code_sku"])
write_df("promo_days", R["promo_days"], index=False)
write_df("unit_economics", R["unit_econ_now"], note="Recent cost regime: Jun–Aug 2026 platform mix and waste; fruit cost = average of last 8 weeks")
write_df("price_sim", R["price_sim"], index=False, note="Constant-elasticity simulation: units = base × (new/old price)^elasticity; margin after commission, fruit, packaging, labor and waste")
write_df("platform", R["platform_pnl"])
write_df("demand_hour", R["demand_hour"], note="Share of each region's daily cups by hour")
write_df("demand_dow", R["demand_dow"].to_frame(), note="Cups per day by weekday (0 = Monday)")
bt = pd.DataFrame({"weekly_wape": R["backtest_wape_weekly"], "weekly_bias": R["backtest_bias_weekly"],
                   "daily_wape": R["backtest_wape_daily"], "daily_error_sd": R["daily_err_sd"]})
write_df("forecast_backtest", bt, note="Forecast = mean of the same weekday over the previous 4 weeks (de-promoted), tested on 1 Jun – 31 Aug 2026")
write_df("newsvendor", R["newsvendor"], note="Critical ratio = lost margin per missed sale ÷ (lost margin + waste cost per wasted cup), at the plan's prices")
for n, a in R["outlook"].items():
    write_df("outlook_" + n.lower().replace(" ", "_").replace("(", "").replace(")", "")[:22], a)
write_df("bridge", R["bridge"], index=False, note="Sep–Nov 2026 gross profit, adding one lever at a time (mid elasticity)")
write_df("kitchen_plan", R["kitchen_plan"], note="Sep–Nov 2026 under the recommended plan")
write_df("inventory_month", R["inventory_month"], index=False)
write_df("inventory_weekly", R["inventory_weekly"], index=False, note="Plan: cups to prep = forecast + waste allowance; fruit spend = prep cups × fruit cost per cup. safety_cups_per_day from the newsvendor buffer (≈0: prep to forecast)")
write_df("forecast_daily_plan", R["forecast_daily_plan"], index=False)
write_df("backcast", R["backcast"], note="Model replayed over Sep 2025 – Aug 2026. model_status_quo_gp vs actual_gp validates the model; plan_op = actual GP + plan uplift − overhead")
write_df("assumptions", R["assumptions"], index=False)
ws = wb["assumptions"]
for col, w in zip("ABC", [32, 100, 60]):
    ws.column_dimensions[col].width = w

wb.save(OUT / "FruitBlend24_analysis.xlsx")
print("saved", len(wb.sheetnames), "sheets")
